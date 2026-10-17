"""Testes com planilhas de exemplo geradas na hora.

Rodar: python -m unittest test_conferencia -v
"""

import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook

import conferencia as c

LOJAS_TXT = """20-256 LOJA MOOCA (SP)
20-1000 LOJA CENTRO (RJ)

20-050
20-073 LOJA PINHEIROS (SP)
"""

CABECALHO_PEDIDO = ["TIPO PEDIDO", "PRODUTO", "QTD.VEN", "PRC.VEN", "TP.OPER", "TES", "CLIENTE", "LOJA",
                    "TABELA", "COND.PG", "VENDEDOR", "ITEM CONTA", "C.CUSTO", "PED.CLIENTE"]
MODELO_REAL = Path(__file__).resolve().parent / c.NOME_MODELO


def criar_planilha(caminho: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendas"
    # A1 não é loja; B1 e E1 repetem 20-256 (duplicada); D1 é desconhecida.
    ws.append(["Data", "  20-256 loja mooca ", "LOJA 20-1000 CENTRO", "20-999 LOJA X", "20-256", None, "Total"])
    ws.append(["linha 2 deve ser ignorada", "20-073", "20-073"])
    ws2 = wb.create_sheet("Extra")
    ws2.append([None, "20-050"])
    wb.save(caminho)


def criar_pedido(caminho: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["EAN", "COD_PP", "COD_ME", "PRODUTO", "20-1000 CENTRO", "20-256 MOOCA", "20-999 X"])
    ws.append(["111", "A1T", 101, "Serum", 0, "ME", 1])           # 20-999 não está no lojas.txt
    ws.append(["222", "B2T", 102, "Creme", 1, " me ", "ME"])
    ws.append(["333", "C3T", 103, "Batom", "X", "Sem estoque", 0])
    ws.append(["444", "D4T", 104, "Base", "1", 1.0, None])
    wb.save(caminho)


def criar_modelo(caminho: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PEDIDO"
    ws.append(CABECALHO_PEDIDO)
    ws["D2"].number_format = "0.00"   # célula vazia com estilo, como no modelo real
    ws.auto_filter.ref = "A1:X3"
    wb.create_sheet("TIPO DE PEDIDO")["A1"] = "TRANSFERENCIA ME"
    wb.save(caminho)


ESPERADO = [
    ("TRANSFERENCIA TESTER 18", "D4T", 1, None, None, None, "FIL018", None, None, None, None, None, None, "20-256"),
    ("TRANSFERENCIA ME", 101, "ME", None, None, None, "FIL018", None, None, None, None, None, None, "20-256"),
    ("TRANSFERENCIA ME", 102, "ME", None, None, None, "FIL018", None, None, None, None, None, None, "20-256"),
    ("TRANSFERENCIA TESTER 18", "B2T", 1, None, None, None, "FIL018", None, None, None, None, None, None, "20-1000"),
    ("TRANSFERENCIA TESTER 18", "D4T", 1, None, None, None, "FIL018", None, None, None, None, None, None, "20-1000"),
]


class TesteConferencia(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.pasta = Path(self.tmp.name)
        self.lojas = self.pasta / "lojas.txt"
        self.lojas.write_text(LOJAS_TXT, encoding="utf-8")
        self.planilha = self.pasta / "Pedido_Teste_exemplo.xlsx"
        criar_planilha(self.planilha)
        self.pedido = self.pasta / "Pedido_Drunk Elephant_ teste outubro.xlsx"
        criar_pedido(self.pedido)
        self.modelo = self.pasta / c.NOME_MODELO
        criar_modelo(self.modelo)

    def tearDown(self):
        self.tmp.cleanup()

    def test_conferencia_das_lojas(self):
        destino, r = c.processar(self.planilha, self.lojas, self.modelo)

        self.assertEqual(destino.name, "PLANILHA DE PEDIDOS - novo.xlsx")
        codigos = [r.lojas[k].codigo for k in r.encontradas]
        self.assertEqual(codigos, ["20-050", "20-256", "20-1000"])  # ordem numérica
        dup = r.duplicadas
        self.assertEqual(list(dup), [(20, 256)])
        self.assertEqual([o.celula for o in dup[(20, 256)]], ["B1", "E1"])
        self.assertEqual([l.codigo for l in r.faltando], ["20-073"])  # só na linha 2
        self.assertEqual([o.valor for o in r.nao_reconhecidas], ["20-999 LOJA X"])
        self.assertEqual(r.pedidos, [])

        # A planilha original não foi alterada.
        self.assertEqual(load_workbook(self.planilha)["Vendas"]["A2"].value, "linha 2 deve ser ignorada")

    def test_pedidos_so_1_e_me(self):
        _, r = c.processar(self.pedido, self.lojas, self.modelo)
        itens = [(r.lojas[it.chave].codigo, it.cod_pp, it.cod_me, it.tipo) for it in r.pedidos]
        self.assertEqual(itens, [
            ("20-256", "D4T", 104, "1"),
            ("20-256", "A1T", 101, "ME"),
            ("20-256", "B2T", 102, "ME"),
            ("20-1000", "B2T", 102, "1"),
            ("20-1000", "D4T", 104, "1"),
        ])

    def test_planilha_de_pedidos_preenchida(self):
        destino, _ = c.processar(self.pedido, self.lojas, self.modelo)

        wb = load_workbook(destino)
        self.assertEqual(wb.sheetnames, ["PEDIDO", "TIPO DE PEDIDO"])
        ws = wb["PEDIDO"]
        linhas = list(ws.iter_rows(max_col=14, values_only=True))
        self.assertEqual(list(linhas[0]), CABECALHO_PEDIDO)
        self.assertEqual(linhas[1:], ESPERADO)
        self.assertEqual([ws.cell(row=n, column=24).value for n in range(2, 7)], ["NÃO"] * 5)  # X = DESCONTO
        msg = [ws.cell(row=n, column=17).value for n in range(1, 7)]
        self.assertEqual(msg, [
            None,  # cabeçalho do modelo de teste só vai até a coluna N
            "20-256 TESTER SEPHORA Drunk Elephant",
            "20-256 TRANSFORMAR TESTER SEPHORA Drunk Elephant",
            "20-256 TRANSFORMAR TESTER SEPHORA Drunk Elephant",
            "20-1000 TESTER SEPHORA Drunk Elephant",
            "20-1000 TESTER SEPHORA Drunk Elephant",
        ])
        self.assertEqual(ws["D2"].number_format, "0.00")   # estilo do modelo mantido
        self.assertEqual(ws.auto_filter.ref, "A1:X6")      # filtro cobre as linhas novas
        self.assertEqual(wb["TIPO DE PEDIDO"]["A1"].value, "TRANSFERENCIA ME")

        # Só a aba PEDIDO e o workbook.xml mudam; o resto é cópia byte a byte do modelo.
        with zipfile.ZipFile(self.modelo) as a, zipfile.ZipFile(destino) as b:
            self.assertEqual(a.namelist(), b.namelist())
            mudou = [n for n in a.namelist() if a.read(n) != b.read(n)]
        self.assertIn("xl/worksheets/sheet1.xml", mudou)
        self.assertLessEqual(set(mudou), {"xl/workbook.xml", "xl/worksheets/sheet1.xml"})

    @unittest.skipUnless(MODELO_REAL.is_file(), "modelo real não está na pasta")
    def test_modelo_real_mantem_listas_suspensas(self):
        destino, _ = c.processar(self.pedido, self.lojas, MODELO_REAL)
        with zipfile.ZipFile(destino) as z:
            self.assertIn(b"x14:dataValidation ", z.read("xl/worksheets/sheet1.xml"))
        ws = load_workbook(destino)["PEDIDO"]
        self.assertEqual(list(ws.iter_rows(min_row=2, max_row=6, max_col=14, values_only=True)), ESPERADO)

    def test_marca_pelo_nome(self):
        self.assertEqual(c.extrair_marca(Path("Pedido_NARS_2026-09-04..2026-10-03.xlsx")), "NARS")
        self.assertEqual(c.extrair_marca(Path("Pedido_Drunk Elephant_ teste outubro - sem formulas.xlsx")),
                         "Drunk Elephant")
        self.assertEqual(c.extrair_marca(Path("pedido_Shiseido_x.xlsx")), "Shiseido")
        for ruim in ("PLANILHA DE PEDIDOS.xlsx", "Pedido_NARS.xlsx", "Pedido__x.xlsx"):
            with self.assertRaises(c.ErroConferencia, msg=ruim):
                c.extrair_marca(Path(ruim))

    def test_pasta_de_destino(self):
        pasta = self.pasta / "saida"
        pasta.mkdir()
        etapas = []
        destino, _ = c.processar(self.pedido, self.lojas, self.modelo, pasta_saida=pasta, progresso=etapas.append)
        self.assertEqual(destino, pasta / c.NOME_SAIDA)
        self.assertTrue(destino.exists())
        self.assertEqual(len(etapas), 5)  # a interface usa esse total para a barra de progresso
        with self.assertRaises(c.ErroConferencia):
            c.processar(self.pedido, self.lojas, self.modelo, pasta_saida=self.pasta / "nao_existe")

    def test_agrupar_lojas_em_ordem_numerica(self):
        chaves = [(20, 1000), (20, 50), (20, 256), (20, 50), (20, 53), (20, 54), (20, 56), (20, 57), (20, 58)]
        self.assertEqual(c.agrupar_lojas(chaves), [
            [(20, 50), (20, 53), (20, 54), (20, 56), (20, 57), (20, 58), (20, 256)],
            [(20, 1000)],
        ])

    def test_lote_divide_em_grupos_de_7(self):
        codigos = ["20-050", "20-053", "20-054", "20-056", "20-057", "20-058", "20-059", "20-070", "20-1000"]
        self.lojas.write_text("\n".join(codigos), encoding="utf-8")
        # Ordem embaralhada no cabeçalho; o lote tem que ordenar numericamente.
        cabecalho = ["COD_PP", "COD_ME", "20-1000 X", "20-059 X", "20-050 X", "20-070 X",
                     "20-053 X", "20-054 X", "20-056 X", "20-057 X", "20-058 X"]
        planilhas = []
        for marca, valor in (("NARS", 1), ("Shiseido", "ME")):
            arq = self.pasta / f"Pedido_{marca}_teste.xlsx"
            wb = Workbook()
            wb.active.append(cabecalho)
            wb.active.append([f"{marca}-PP", f"{marca}-ME"] + [valor] * 9)
            wb.save(arq)
            planilhas.append(arq)
        saida = self.pasta / "saida"
        saida.mkdir()

        r = c.processar_lote(planilhas, self.lojas, self.modelo, pasta_saida=saida)

        self.assertEqual([a.destino.name for a in r.arquivos],
                         ["PLANILHA DE PEDIDOS - (1).xlsx", "PLANILHA DE PEDIDOS - (2).xlsx"])
        self.assertEqual([l.codigo for l in r.arquivos[0].lojas], codigos[:7])
        self.assertEqual([l.codigo for l in r.arquivos[1].lojas], ["20-070", "20-1000"])

        ws = load_workbook(saida / "PLANILHA DE PEDIDOS - (2).xlsx")["PEDIDO"]
        linhas = [(row[0], row[1], row[2], row[13], row[16])
                  for row in ws.iter_rows(min_row=2, max_col=17, values_only=True)]
        # Cada loja traz os pedidos das duas planilhas, na ordem em que foram anexadas.
        self.assertEqual(linhas, [
            ("TRANSFERENCIA TESTER 18", "NARS-PP", 1, "20-070", "20-070 TESTER SEPHORA NARS"),
            ("TRANSFERENCIA ME", "Shiseido-ME", "ME", "20-070", "20-070 TRANSFORMAR TESTER SEPHORA Shiseido"),
            ("TRANSFERENCIA TESTER 18", "NARS-PP", 1, "20-1000", "20-1000 TESTER SEPHORA NARS"),
            ("TRANSFERENCIA ME", "Shiseido-ME", "ME", "20-1000", "20-1000 TRANSFORMAR TESTER SEPHORA Shiseido"),
        ])
        self.assertEqual(len(r.arquivos[0].itens), 14)  # 7 lojas × 2 planilhas

    def test_lote_erro_diz_qual_planilha(self):
        ruim = self.pasta / "sem_padrao.xlsx"
        criar_pedido(ruim)
        with self.assertRaises(c.ErroConferencia) as ctx:
            c.processar_lote([self.pedido, ruim], self.lojas, self.modelo)
        self.assertIn("sem_padrao.xlsx", str(ctx.exception))

    def test_ped_cliente_ganha_sufixo_a_cada_99_msg_nota(self):
        lojas = c.ler_lojas(self.lojas)
        tester = [c.ItemPedido((20, 50), "", "PP", "ME", "1", "NARS") for _ in range(250)]
        me = [c.ItemPedido((20, 50), "", "PP", "ME", "ME", "NARS") for _ in range(5)]
        outra_marca = [c.ItemPedido((20, 50), "", "PP", "ME", "1", "Shiseido") for _ in range(3)]
        linhas = c.linhas_do_pedido(tester + me + outra_marca, lojas)

        ped = [l["N"] for l in linhas]
        self.assertEqual(ped[:250], ["20-050"] * 99 + ["20-0501"] * 99 + ["20-0502"] * 52)
        # Cada MSG NOTA conta separado: ME e outra marca começam do zero.
        self.assertEqual(ped[250:], ["20-050"] * 8)
        # A MSG NOTA não muda.
        self.assertEqual({l["Q"] for l in linhas[:250]}, {"20-050 TESTER SEPHORA NARS"})

    def _criar_precos(self) -> Path:
        arq = self.pasta / "precos.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(["CODIGO", "TIPO", "GRUPO", "DESCRICAO", "U.M.", "FILIAL", "ARMAZEM", "SALDO",
                   "EMPENHO", "DISPONIVEL", "VALOR UNITARIO", "TOTAL"])
        ws.append(["D4T", "PP", "", "", "", "", "", 0, 0, 0, 31.24, 0])
        ws.append([101, "ME", "", "", "", "", "", 0, 0, 0, 37.82123, 0])    # código numérico
        ws.append(["B2T", "PP", "", "", "", "", "", 0, 0, 0, 10, 0])
        wb.save(arq)
        return arq

    def test_ler_precos(self):
        precos = c.ler_precos(self._criar_precos())
        self.assertEqual(precos, {"D4T": 31.24, "101": 37.82123, "B2T": 10})
        with self.assertRaises(c.ErroConferencia):
            c.ler_precos(self.pasta / "nao_existe.xlsx")

    def test_preco_na_coluna_d_em_reais(self):
        destino, r = c.processar(self.pedido, self.lojas, self.modelo, planilha_precos=self._criar_precos())
        ws = load_workbook(destino)["PEDIDO"]
        self.assertEqual([ws.cell(row=n, column=4).value for n in range(2, 7)],
                         [31.24, 37.82123, None, 10, 31.24])   # 102 (ME) não está na planilha de preços
        self.assertEqual(r.sem_preco, ["102"])
        # Toda a coluna D usa o formato da D2 do modelo (moeda).
        self.assertEqual({ws.cell(row=n, column=4).number_format for n in range(2, 7)}, {"0.00"})

    def test_sem_planilha_de_precos_coluna_d_vazia(self):
        destino, r = c.processar(self.pedido, self.lojas, self.modelo)
        ws = load_workbook(destino)["PEDIDO"]
        self.assertEqual([ws.cell(row=n, column=4).value for n in range(2, 7)], [None] * 5)
        self.assertEqual(r.sem_preco, [])

    def test_lote_com_tamanho_de_grupo_escolhido(self):
        r = c.processar_lote([self.pedido], self.lojas, self.modelo, tamanho_grupo=1,
                             planilha_precos=self._criar_precos())
        self.assertEqual([a.destino.name for a in r.arquivos],
                         ["PLANILHA DE PEDIDOS - (1).xlsx", "PLANILHA DE PEDIDOS - (2).xlsx"])
        self.assertEqual([[l.codigo for l in a.lojas] for a in r.arquivos], [["20-256"], ["20-1000"]])
        self.assertEqual(r.sem_preco, ["102"])
        with self.assertRaises(c.ErroConferencia):
            c.processar_lote([self.pedido], self.lojas, self.modelo, tamanho_grupo=0)

    def test_modelo_ausente(self):
        with self.assertRaises(c.ErroConferencia):
            c.processar(self.pedido, self.lojas, self.pasta / "nao_existe.xlsx")

    def test_lojas_latin1(self):
        self.lojas.write_bytes("20-256 LOJA SÃO PAULO (SP)\n".encode("latin-1"))
        self.assertEqual(c.ler_lojas(self.lojas)[(20, 256)].nome, "LOJA SÃO PAULO (SP)")

    def test_lojas_txt_ausente(self):
        with self.assertRaises(c.ErroConferencia):
            c.ler_lojas(self.pasta / "nao_existe.txt")

    def test_planilha_corrompida(self):
        ruim = self.pasta / "ruim.xlsx"
        ruim.write_bytes(b"isto nao e um xlsx")
        with self.assertRaises(c.ErroConferencia):
            c.ler_cabecalhos(ruim)

    def test_planilha_sem_linha1(self):
        vazia = self.pasta / "vazia.xlsx"
        wb = Workbook()
        wb.active["A2"] = "20-256"
        wb.save(vazia)
        with self.assertRaises(c.ErroConferencia):
            c.ler_cabecalhos(vazia)


if __name__ == "__main__":
    unittest.main()
