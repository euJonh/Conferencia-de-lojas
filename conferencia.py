"""Conferência de lojas na linha 1 de uma planilha .xlsx.

Lê o cabeçalho (linha 1) de todas as abas da planilha escolhida pelo usuário,
compara com a lista oficial do lojas.txt e pega as células com 1 ou ME abaixo de
cada loja (0, X e outros valores são ignorados). Gera "PLANILHA DE PEDIDOS - novo.xlsx",
cópia idêntica do modelo "PLANILHA DE PEDIDOS.xlsx", com a aba PEDIDO preenchida:
A = TIPO PEDIDO, B = PRODUTO (COD_PP ou COD_ME), C = QTD.VEN (1 ou ME),
D = PRC.VEN (VALOR UNITARIO da planilha de preços, pelo código da coluna B),
G = CLIENTE (FIL018), N = PED.CLIENTE (código da loja), Q = MSG NOTA (código +
TESTER SEPHORA ou TRANSFORMAR TESTER SEPHORA + marca do nome Pedido_<MARCA>_...) e
X = DESCONTO (NÃO). A conferência das lojas (encontradas, duplicadas,
faltando...) aparece só na tela de resultado.

Em lote (processar_lote), várias planilhas são juntadas e divididas em arquivos
"PLANILHA DE PEDIDOS - (1).xlsx", "(2)"... com N lojas cada (padrão 7), em ordem numérica.
"""

from __future__ import annotations

import re
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

NOME_LOJAS_TXT = "lojas.txt"
NOME_MODELO = "PLANILHA DE PEDIDOS.xlsx"
NOME_SAIDA = "PLANILHA DE PEDIDOS - novo.xlsx"

# Código de loja: NN-NNN (a parte depois do hífen pode ter 3 ou mais dígitos, ex.: 20-1000).
# Os lookarounds impedem casar pedaços de números maiores (ex.: "120-2560").
PADRAO_CODIGO = re.compile(r"(?<![\d-])(\d{2})-(\d{3,})(?![\d-])")
# Linha do lojas.txt: "20-256" ou "20-256 LOJA MOOCA (SP)".
PADRAO_LINHA_LOJAS = re.compile(r"^(\d+)-(\d+)\s*(.*)$")


class ErroConferencia(Exception):
    """Erro com mensagem amigável, pronta para exibir ao usuário."""


@dataclass
class Loja:
    codigo: str
    nome: str

    @property
    def chave(self) -> tuple[int, int]:
        return chave_codigo(self.codigo)


@dataclass
class Ocorrencia:
    aba: str
    celula: str  # ex.: "C1"
    valor: str


@dataclass
class ItemPedido:
    """Célula com 1 ou ME abaixo de uma loja conhecida."""
    chave: tuple[int, int]
    loja: str      # texto do cabeçalho da coluna da loja
    cod_pp: object
    cod_me: object
    tipo: str      # "1" ou "ME"
    marca: str = ""  # marca da planilha de onde o item veio (vai na MSG NOTA)


@dataclass
class Resultado:
    lojas: dict[tuple[int, int], Loja]
    encontradas: dict[tuple[int, int], list[Ocorrencia]] = field(default_factory=dict)
    nao_reconhecidas: list[Ocorrencia] = field(default_factory=list)
    pedidos: list[ItemPedido] = field(default_factory=list)
    marca: str = ""  # do nome da planilha: Pedido_<MARCA>_...
    sem_preco: list[str] = field(default_factory=list)  # códigos sem preço na planilha de preços

    @property
    def duplicadas(self) -> dict[tuple[int, int], list[Ocorrencia]]:
        return {k: v for k, v in self.encontradas.items() if len(v) > 1}

    @property
    def faltando(self) -> list[Loja]:
        return [self.lojas[k] for k in sorted(self.lojas) if k not in self.encontradas]


def chave_codigo(codigo: str) -> tuple[int, int]:
    """Chave numérica para ordenar/comparar códigos: '20-256' -> (20, 256)."""
    antes, depois = codigo.strip().split("-", 1)
    return int(antes), int(depois)


def pasta_do_programa() -> Path:
    """Pasta do .exe quando empacotado, ou do script quando rodando com Python."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


# ---------------------------------------------------------------------------
# Leitura do lojas.txt
# ---------------------------------------------------------------------------

def ler_lojas(caminho: Path) -> dict[tuple[int, int], Loja]:
    if not caminho.is_file():
        raise ErroConferencia(
            f"Arquivo {NOME_LOJAS_TXT} não encontrado.\n\n"
            f"Coloque o {NOME_LOJAS_TXT} na mesma pasta do programa:\n{caminho.parent}"
        )

    try:
        texto = caminho.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        texto = caminho.read_text(encoding="latin-1")

    lojas: dict[tuple[int, int], Loja] = {}
    for linha in texto.splitlines():
        linha = linha.strip()
        if not linha:
            continue
        m = PADRAO_LINHA_LOJAS.match(linha)
        if not m:
            continue  # linha fora do formato esperado
        codigo = f"{m.group(1)}-{m.group(2)}"
        loja = Loja(codigo=codigo, nome=m.group(3).strip())
        lojas.setdefault(loja.chave, loja)

    if not lojas:
        raise ErroConferencia(
            f"Nenhuma loja válida encontrada no {NOME_LOJAS_TXT}.\n\n"
            "Formato esperado, uma por linha: 20-256 ou 20-256 LOJA MOOCA (SP)"
        )
    return lojas


# ---------------------------------------------------------------------------
# Leitura da planilha
# ---------------------------------------------------------------------------

def ler_planilha(caminho: Path) -> dict[str, list[tuple]]:
    """Retorna {aba: [linha1, linha2, ...]} com os valores de todas as linhas."""
    try:
        wb = load_workbook(caminho, read_only=True, data_only=True)
    except PermissionError:
        raise ErroConferencia(
            "Não foi possível abrir a planilha porque ela está em uso.\n\n"
            "Feche o arquivo no Excel (ou em outro programa) e tente novamente."
        )
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError):
        raise ErroConferencia(
            "Não foi possível ler a planilha. O arquivo pode estar corrompido "
            "ou não ser um .xlsx válido."
        )

    try:
        return {ws.title: list(ws.iter_rows(values_only=True)) for ws in wb.worksheets}
    finally:
        wb.close()


def extrair_cabecalhos(linhas_por_aba: dict[str, list[tuple]]) -> dict[str, list[tuple[str, str]]]:
    """Retorna {aba: [(celula, valor), ...]} só com as células preenchidas da linha 1."""
    cabecalhos: dict[str, list[tuple[str, str]]] = {}
    for aba, linhas in linhas_por_aba.items():
        celulas = []
        for col, valor in enumerate(linhas[0] if linhas else (), start=1):
            valor = "" if valor is None else str(valor).strip()
            if valor:
                celulas.append((f"{get_column_letter(col)}1", valor))
        cabecalhos[aba] = celulas

    if not any(cabecalhos.values()):
        raise ErroConferencia("A planilha não tem dados na linha 1 (cabeçalho) de nenhuma aba.")
    return cabecalhos


def ler_cabecalhos(caminho: Path) -> dict[str, list[tuple[str, str]]]:
    return extrair_cabecalhos(ler_planilha(caminho))


# ---------------------------------------------------------------------------
# Planilha de preços (CODIGO -> VALOR UNITARIO)
# ---------------------------------------------------------------------------

COLUNA_CODIGO_PRECO = "CODIGO"            # coluna A da planilha de preços
COLUNA_VALOR_PRECO = "VALOR UNITARIO"     # coluna K da planilha de preços


def normalizar_codigo(valor) -> str:
    """Chave para casar códigos de produto: 42800352, 42800352.0 e " 42800352 " viram "42800352"."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip().upper()


def ler_precos(caminho: Path) -> dict[str, float]:
    """{código do produto: valor unitário}. Colunas localizadas pelo nome do cabeçalho
    (CODIGO e VALOR UNITARIO); sem cabeçalho reconhecível, usa A e K."""
    if not caminho.is_file():
        raise ErroConferencia(
            f"Planilha de preços não encontrada:\n{caminho}\n\n"
            "Anexe a planilha de preços novamente em Configurações."
        )
    linhas_por_aba = ler_planilha(caminho)
    for linhas in linhas_por_aba.values():
        if not linhas:
            continue
        titulos = [normalizar_codigo(t) if t is not None else "" for t in linhas[0]]
        if COLUNA_CODIGO_PRECO in titulos and COLUNA_VALOR_PRECO in titulos:
            i_cod, i_val = titulos.index(COLUNA_CODIGO_PRECO), titulos.index(COLUNA_VALOR_PRECO)
            break
    else:
        linhas, i_cod, i_val = next(iter(linhas_por_aba.values())), 0, 10

    precos: dict[str, float] = {}
    for linha in linhas[1:]:
        if len(linha) <= max(i_cod, i_val) or linha[i_cod] is None:
            continue
        valor = linha[i_val]
        if isinstance(valor, (int, float)) and not isinstance(valor, bool):
            precos.setdefault(normalizar_codigo(linha[i_cod]), valor)
    if not precos:
        raise ErroConferencia(
            "Nenhum preço encontrado na planilha de preços.\n\n"
            f"Ela precisa das colunas {COLUNA_CODIGO_PRECO} (A) e {COLUNA_VALOR_PRECO} (K)."
        )
    return precos


def codigo_produto(it: ItemPedido):
    """Código que vai na coluna B: COD_PP para 1, COD_ME para ME."""
    return it.cod_pp if it.tipo == "1" else it.cod_me


def codigos_sem_preco(itens: list[ItemPedido], precos: dict[str, float] | None) -> list[str]:
    if precos is None:
        return []
    faltam = {normalizar_codigo(codigo_produto(it)) for it in itens} - precos.keys()
    return sorted(faltam)


# ---------------------------------------------------------------------------
# Pedidos (células 1 e ME abaixo de cada loja)
# ---------------------------------------------------------------------------

def tipo_pedido(valor) -> str | None:
    """'1' ou 'ME' quando a célula pede o produto; None para 0, X e qualquer outro valor."""
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return "1" if valor == 1 else None
    if isinstance(valor, str):
        texto = valor.strip().upper()
        return texto if texto in ("1", "ME") else None
    return None


def extrair_pedidos(lojas: dict[tuple[int, int], Loja],
                    linhas_por_aba: dict[str, list[tuple]]) -> list[ItemPedido]:
    itens: list[ItemPedido] = []
    for linhas in linhas_por_aba.values():
        if not linhas:
            continue
        cabecalho = linhas[0]

        # Colunas COD_PP e COD_ME pelo nome do cabeçalho; ausentes ficam vazias.
        idx_produto: dict[str, int] = {}
        colunas_loja: list[tuple[int, tuple[int, int], str]] = []
        for i, titulo in enumerate(cabecalho):
            if titulo is None:
                continue
            texto = str(titulo).strip()
            idx_produto.setdefault(texto.upper(), i)
            for m in PADRAO_CODIGO.finditer(texto):
                chave = (int(m.group(1)), int(m.group(2)))
                if chave in lojas:
                    colunas_loja.append((i, chave, texto))
                    break

        def pegar(linha: tuple, nome: str):
            i = idx_produto.get(nome)
            return linha[i] if i is not None and i < len(linha) else None

        for linha in linhas[1:]:
            for i, chave, texto_loja in colunas_loja:
                tipo = tipo_pedido(linha[i]) if i < len(linha) else None
                if tipo:
                    itens.append(ItemPedido(
                        chave=chave, loja=texto_loja,
                        cod_pp=pegar(linha, "COD_PP"), cod_me=pegar(linha, "COD_ME"), tipo=tipo,
                    ))

    # Por loja, depois 1 antes de ME; dentro disso mantém a ordem da planilha.
    itens.sort(key=lambda it: (it.chave, it.tipo != "1"))
    return itens


# ---------------------------------------------------------------------------
# Comparação
# ---------------------------------------------------------------------------

def comparar(lojas: dict[tuple[int, int], Loja],
             cabecalhos: dict[str, list[tuple[str, str]]]) -> Resultado:
    resultado = Resultado(lojas=lojas)

    for aba, celulas in cabecalhos.items():
        for celula, valor in celulas:
            ocorrencia = Ocorrencia(aba=aba, celula=celula, valor=valor)
            # Uma célula pode citar mais de um código; cada um conta separado.
            for m in PADRAO_CODIGO.finditer(valor):
                chave = (int(m.group(1)), int(m.group(2)))
                if chave in lojas:
                    resultado.encontradas.setdefault(chave, []).append(ocorrencia)
                elif ocorrencia not in resultado.nao_reconhecidas:
                    resultado.nao_reconhecidas.append(ocorrencia)

    resultado.encontradas = dict(sorted(resultado.encontradas.items()))
    return resultado


# ---------------------------------------------------------------------------
# Planilha de pedidos
# ---------------------------------------------------------------------------
# O modelo tem listas suspensas (validação de dados do Excel 2010+) que o openpyxl
# apaga ao salvar. Por isso a cópia é feita arquivo a arquivo dentro do .xlsx e só o
# XML da aba PEDIDO (e o intervalo do filtro) é editado; o resto fica idêntico ao modelo.

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

ABA_PEDIDO = "PEDIDO"
CLIENTE = "FIL018"
DESCONTO = "NÃO"
TIPO_PEDIDO = {"1": "TRANSFERENCIA TESTER 18", "ME": "TRANSFERENCIA ME"}
MSG_NOTA = {"1": "TESTER SEPHORA", "ME": "TRANSFORMAR TESTER SEPHORA"}
LIMITE_MSG_NOTA = 99  # linhas por MSG NOTA antes de o PED.CLIENTE ganhar sufixo (1, 2...)

# Nome da planilha do usuário: Pedido_<NOME DA MARCA>_<resto>.xlsx
PADRAO_NOME_PLANILHA = re.compile(r"^\s*Pedido_(.+?)_", re.I)

PADRAO_LINHA_XML = re.compile(r'<row\b[^>]*?\br="(\d+)"[^>]*?(?:/>|>(.*?)</row>)', re.S)
PADRAO_CELULA_XML = re.compile(r'<c\b[^>]*?\br="([A-Z]+)\d+"[^>]*?(?:/>|>.*?</c>)', re.S)
PADRAO_ESTILO = re.compile(r'\bs="(\d+)"')


def extrair_marca(planilha: Path) -> str:
    """'Pedido_NARS_2026-09-04..2026-10-03.xlsx' -> 'NARS'."""
    m = PADRAO_NOME_PLANILHA.match(planilha.stem)
    marca = m.group(1).strip() if m else ""
    if not marca:
        raise ErroConferencia(
            "Não foi possível identificar a marca pelo nome da planilha.\n\n"
            f"Nome recebido: {planilha.name}\n"
            "Formato esperado: Pedido_NOME DA MARCA_....xlsx"
        )
    return marca


def linhas_do_pedido(itens: list[ItemPedido], lojas: dict[tuple[int, int], Loja],
                     precos: dict[str, float] | None = None) -> list[dict[str, object]]:
    """Uma linha da aba PEDIDO por item: {coluna: valor} para A, B, C, D, G, N, Q e X.

    PRC.VEN (D) vem da planilha de preços pelo código da coluna B; sem preço, fica vazio.

    PED.CLIENTE: a cada LIMITE_MSG_NOTA linhas com a mesma MSG NOTA, o código da loja
    ganha um sufixo: linhas 1-99 "20-050", 100-198 "20-0501", 199-297 "20-0502"...
    """
    linhas = []
    repeticoes: dict[str, int] = {}
    for it in itens:
        codigo = lojas[it.chave].codigo
        msg = f"{codigo} {MSG_NOTA[it.tipo]} {it.marca}"
        bloco = repeticoes.get(msg, 0) // LIMITE_MSG_NOTA
        repeticoes[msg] = repeticoes.get(msg, 0) + 1
        linhas.append({
            "A": TIPO_PEDIDO[it.tipo],
            "B": codigo_produto(it),
            "C": 1 if it.tipo == "1" else "ME",
            "D": precos.get(normalizar_codigo(codigo_produto(it))) if precos else None,
            "G": CLIENTE,
            "N": f"{codigo}{bloco}" if bloco else codigo,
            "Q": msg,
            "X": DESCONTO,
        })
    return linhas


def _indice_coluna(letras: str) -> int:
    n = 0
    for ch in letras:
        n = n * 26 + ord(ch) - 64
    return n


def _xml_celula(ref: str, valor, estilo: str | None) -> str:
    s = f' s="{estilo}"' if estilo else ""
    if valor is None or valor == "":
        return f'<c r="{ref}"{s}/>'
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return f'<c r="{ref}"{s}><v>{valor}</v></c>'
    texto = escape(str(valor))
    return f'<c r="{ref}"{s} t="inlineStr"><is><t xml:space="preserve">{texto}</t></is></c>'


def _arquivo_da_aba(arquivos: dict[str, bytes], nome_aba: str) -> str:
    """Caminho dentro do .xlsx (ex.: xl/worksheets/sheet1.xml) da aba com esse nome."""
    wb = ET.fromstring(arquivos["xl/workbook.xml"])
    rid = None
    for sh in wb.iter(f"{{{NS_MAIN}}}sheet"):
        if sh.get("name", "").strip().upper() == nome_aba:
            rid = sh.get(f"{{{NS_REL}}}id")
            break
    if rid is None:
        raise ErroConferencia(f"O modelo {NOME_MODELO} não tem a aba {nome_aba}.")
    rels = ET.fromstring(arquivos["xl/_rels/workbook.xml.rels"])
    for rel in rels:
        if rel.get("Id") == rid:
            alvo = rel.get("Target", "")
            return alvo.lstrip("/") if alvo.startswith("/") else f"xl/{alvo}"
    raise ErroConferencia(f"O modelo {NOME_MODELO} está corrompido (aba {nome_aba} sem arquivo).")


def preencher_aba_xml(xml: str, linhas: list[dict[str, object]],
                      estilo_linha2: tuple[str, ...] = ()) -> tuple[str, int]:
    """Escreve as linhas a partir da linha 2 do sheetData. Retorna (xml, última linha).

    Cada célula mantém o estilo que já tinha no modelo; sem estilo, herda o da mesma
    coluna na linha 2. As colunas em `estilo_linha2` sempre usam o estilo da linha 2
    (ex.: D, para todo preço sair no formato R$ do modelo)."""
    ini = xml.index(">", xml.index("<sheetData")) + 1
    if xml[ini - 2] == "/":  # <sheetData/> vazio
        xml = xml[:ini - 2] + "></sheetData>" + xml[ini:]
        ini -= 1
    fim = xml.index("</sheetData>")

    # Linhas do modelo: {número: (tag de abertura sem o ">", {coluna: xml da célula})}
    existentes: dict[int, tuple[str, dict[str, str]]] = {}
    for m in PADRAO_LINHA_XML.finditer(xml[ini:fim]):
        tag = re.match(r"<row\b[^>]*?(?=/?>)", m.group(0)).group(0)
        celulas = {c.group(1): c.group(0) for c in PADRAO_CELULA_XML.finditer(m.group(2) or "")}
        existentes[int(m.group(1))] = (tag, celulas)

    def estilo_de(celula_xml: str | None) -> str | None:
        m = PADRAO_ESTILO.search(celula_xml.split(">", 1)[0]) if celula_xml else None
        return m.group(1) if m else None

    estilos_linha2 = {col: estilo_de(cel) for col, cel in existentes.get(2, ("", {}))[1].items()}

    for n, dados in enumerate(linhas, start=2):
        tag, celulas = existentes.get(n, (f'<row r="{n}"', {}))
        tag = re.sub(r'\sspans="[^"]*"', "", tag)  # spans é só dica; evita valor desatualizado
        for col, valor in dados.items():
            if col in estilo_linha2:
                estilo = estilos_linha2.get(col)
            else:
                estilo = estilo_de(celulas.get(col)) or estilos_linha2.get(col)
            celulas[col] = _xml_celula(f"{col}{n}", valor, estilo)
        existentes[n] = (tag, celulas)

    partes = []
    for n in sorted(existentes):
        tag, celulas = existentes[n]
        corpo = "".join(celulas[c] for c in sorted(celulas, key=_indice_coluna))
        partes.append(f"{tag}>{corpo}</row>" if corpo else f"{tag}/>")

    ultima = max([1, *existentes])
    return xml[:ini] + "".join(partes) + xml[fim:], ultima


def _ampliar_intervalo(xml: str, padrao: str, ultima: int) -> str:
    """Faz refs como A1:X277 cobrirem até a última linha preenchida (nunca diminui)."""
    return re.sub(padrao, lambda m: f"{m.group(1)}{max(int(m.group(2)), ultima)}", xml)


def gerar_planilha_pedidos(itens: list[ItemPedido], lojas: dict[tuple[int, int], Loja],
                           modelo: Path, destino: Path, precos: dict[str, float] | None = None) -> None:
    if not modelo.is_file():
        raise ErroConferencia(
            f"Modelo {NOME_MODELO} não encontrado.\n\n"
            f"Coloque o {NOME_MODELO} na mesma pasta do programa:\n{modelo.parent}"
        )
    try:
        with zipfile.ZipFile(modelo) as zin:
            infos = zin.infolist()
            arquivos = {i.filename: zin.read(i.filename) for i in infos}
    except PermissionError:
        raise ErroConferencia(f"O modelo {NOME_MODELO} está aberto em outro programa. Feche-o e tente novamente.")
    except (zipfile.BadZipFile, OSError):
        raise ErroConferencia(f"O modelo {NOME_MODELO} está corrompido ou não é um .xlsx válido.")

    aba = _arquivo_da_aba(arquivos, ABA_PEDIDO)
    xml, ultima = preencher_aba_xml(arquivos[aba].decode("utf-8"), linhas_do_pedido(itens, lojas, precos),
                                    estilo_linha2=("D",))
    xml = _ampliar_intervalo(xml, r'(<dimension ref="[A-Z]+\d+:[A-Z]+)(\d+)', ultima)
    xml = _ampliar_intervalo(xml, r'(<autoFilter ref="[A-Z]+\d+:[A-Z]+)(\d+)', ultima)
    arquivos[aba] = xml.encode("utf-8")
    # Intervalo oculto do filtro (_FilterDatabase) também precisa cobrir as linhas novas.
    wb_xml = arquivos["xl/workbook.xml"].decode("utf-8")
    wb_xml = _ampliar_intervalo(wb_xml, rf"(>{ABA_PEDIDO}!\$[A-Z]+\$\d+:\$[A-Z]+\$)(\d+)", ultima)
    arquivos["xl/workbook.xml"] = wb_xml.encode("utf-8")

    try:
        with zipfile.ZipFile(destino, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in infos:
                zout.writestr(info, arquivos[info.filename])
    except PermissionError:
        raise ErroConferencia(
            f"Não foi possível salvar a planilha:\n{destino}\n\n"
            "Se ela estiver aberta no Excel, feche-a e tente novamente."
        )


def caminho_saida(planilha: Path, pasta_saida: Path | None = None) -> Path:
    """Arquivo de saída na pasta escolhida nas configurações ou, sem ela, junto da planilha."""
    return (pasta_saida or planilha.parent) / NOME_SAIDA


def _checar_pasta_saida(pasta_saida: Path | None) -> None:
    if pasta_saida is not None and not pasta_saida.is_dir():
        raise ErroConferencia(
            f"A pasta de destino não existe:\n{pasta_saida}\n\n"
            "Escolha outra pasta em Configurações."
        )


def _ler_pedido(planilha: Path, lojas: dict[tuple[int, int], Loja]) -> Resultado:
    """Conferência + itens 1/ME de uma planilha, com a marca do nome do arquivo."""
    marca = extrair_marca(planilha)
    linhas_por_aba = ler_planilha(planilha)
    resultado = comparar(lojas, extrair_cabecalhos(linhas_por_aba))
    resultado.marca = marca
    resultado.pedidos = extrair_pedidos(lojas, linhas_por_aba)
    for it in resultado.pedidos:
        it.marca = marca
    return resultado


def processar(planilha: Path, lojas_txt: Path, modelo: Path,
              pasta_saida: Path | None = None,
              progresso: Callable[[str], None] | None = None,
              planilha_precos: Path | None = None) -> tuple[Path, Resultado]:
    """Fluxo completo. `progresso`, se informado, recebe o texto de cada etapa."""
    avisar = progresso or (lambda _texto: None)
    _checar_pasta_saida(pasta_saida)

    avisar("Identificando a marca")
    extrair_marca(planilha)  # falha cedo se o nome estiver fora do padrão
    avisar("Lendo o lojas.txt e os preços")
    lojas = ler_lojas(lojas_txt)
    precos = ler_precos(planilha_precos) if planilha_precos else None
    avisar("Lendo a planilha")
    avisar("Conferindo as lojas")
    resultado = _ler_pedido(planilha, lojas)
    avisar("Gerando a planilha de pedidos")
    destino = caminho_saida(planilha, pasta_saida)
    gerar_planilha_pedidos(resultado.pedidos, lojas, modelo, destino, precos)
    resultado.sem_preco = codigos_sem_preco(resultado.pedidos, precos)
    return destino, resultado


# ---------------------------------------------------------------------------
# Geração em lote
# ---------------------------------------------------------------------------
# Várias planilhas anexadas: os pedidos de todas são juntados e divididos em
# arquivos de TAMANHO_GRUPO lojas, em ordem numérica de código. Cada arquivo leva
# os pedidos dessas lojas vindos de todas as planilhas.

TAMANHO_GRUPO = 7


def nome_saida_lote(numero: int) -> str:
    return f"PLANILHA DE PEDIDOS - ({numero}).xlsx"


@dataclass
class ArquivoLote:
    destino: Path
    lojas: list[Loja]
    itens: list[ItemPedido]


@dataclass
class ResultadoLote:
    planilhas: list[tuple[Path, Resultado]]
    arquivos: list[ArquivoLote]
    sem_preco: list[str] = field(default_factory=list)  # códigos sem preço na planilha de preços


def agrupar_lojas(chaves, tamanho: int = TAMANHO_GRUPO) -> list[list[tuple[int, int]]]:
    """Códigos em ordem numérica, divididos em grupos de `tamanho`."""
    ordenadas = sorted(set(chaves))
    return [ordenadas[i:i + tamanho] for i in range(0, len(ordenadas), tamanho)]


def processar_lote(planilhas: list[Path], lojas_txt: Path, modelo: Path,
                   pasta_saida: Path | None = None,
                   progresso: Callable[[str, float], None] | None = None,
                   tamanho_grupo: int = TAMANHO_GRUPO,
                   planilha_precos: Path | None = None) -> ResultadoLote:
    """`progresso`, se informado, recebe (texto da etapa, fração concluída de 0 a 1)."""
    avisar = progresso or (lambda _texto, _fracao: None)
    if not planilhas:
        raise ErroConferencia("Nenhuma planilha anexada.")
    _checar_pasta_saida(pasta_saida)

    if tamanho_grupo < 1:
        raise ErroConferencia("O número de lojas por arquivo precisa ser pelo menos 1.")
    avisar("Lendo o lojas.txt e os preços", 0.0)
    lojas = ler_lojas(lojas_txt)
    precos = ler_precos(planilha_precos) if planilha_precos else None

    lidas: list[tuple[Path, Resultado]] = []
    for i, planilha in enumerate(planilhas):
        avisar(f"Lendo {planilha.name} ({i + 1} de {len(planilhas)})", 0.6 * i / len(planilhas))
        try:
            lidas.append((planilha, _ler_pedido(planilha, lojas)))
        except ErroConferencia as e:
            raise ErroConferencia(f"Problema na planilha {planilha.name}:\n\n{e}") from None

    # Ordem final em cada arquivo: loja, depois a ordem em que as planilhas foram
    # anexadas, depois 1 antes de ME; dentro disso mantém a ordem da planilha.
    todos = [(ordem, it) for ordem, (_, r) in enumerate(lidas) for it in r.pedidos]
    if not todos:
        raise ErroConferencia("Nenhuma célula com 1 ou ME foi encontrada nas planilhas anexadas.")
    todos.sort(key=lambda par: (par[1].chave, par[0], par[1].tipo != "1"))

    pasta = pasta_saida or planilhas[0].parent
    grupos = agrupar_lojas((it.chave for _, it in todos), tamanho_grupo)
    arquivos = []
    for n, grupo in enumerate(grupos, start=1):
        avisar(f"Gerando {nome_saida_lote(n)} ({n} de {len(grupos)})", 0.6 + 0.4 * (n - 1) / len(grupos))
        do_grupo = set(grupo)
        itens = [it for _, it in todos if it.chave in do_grupo]
        destino = pasta / nome_saida_lote(n)
        gerar_planilha_pedidos(itens, lojas, modelo, destino, precos)
        arquivos.append(ArquivoLote(destino=destino, lojas=[lojas[k] for k in grupo], itens=itens))

    return ResultadoLote(planilhas=lidas, arquivos=arquivos,
                         sem_preco=codigos_sem_preco([it for _, it in todos], precos))


def texto_resumo(resultado: Resultado, destino: Path) -> str:
    return (
        f"Marca: {resultado.marca}\n"
        f"Lojas no {NOME_LOJAS_TXT}: {len(resultado.lojas)}\n"
        f"Encontradas: {len(resultado.encontradas)}\n"
        f"Duplicadas: {len(resultado.duplicadas)}\n"
        f"Faltando: {len(resultado.faltando)}\n"
        f"Não reconhecidas: {len(resultado.nao_reconhecidas)}\n"
        f"Pedidos com 1: {sum(it.tipo == '1' for it in resultado.pedidos)}\n"
        f"Pedidos com ME: {sum(it.tipo == 'ME' for it in resultado.pedidos)}\n\n"
        f"Planilha salva em:\n{destino}"
    )


def main() -> None:
    from interface import iniciar  # interface gráfica fica em interface.py
    iniciar()


if __name__ == "__main__":
    main()
